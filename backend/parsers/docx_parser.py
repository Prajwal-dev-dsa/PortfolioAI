from docx import Document


def extract_docx_text(file) -> str:
    """
    Extract text from DOCX paragraphs and tables.
    """

    document = Document(file)

    sections: list[str] = []

    # Extract normal paragraphs.
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            sections.append(text)

    # Extract table content too.
    # Job descriptions sometimes contain requirements
    # inside tables.
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            row_text = " | ".join(
                cell for cell in cells
                if cell
            )

            if row_text:
                sections.append(row_text)

    return "\n\n".join(sections).strip()